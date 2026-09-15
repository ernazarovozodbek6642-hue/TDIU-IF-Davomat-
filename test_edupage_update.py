import io
import unittest
from pathlib import Path
from unittest.mock import patch, AsyncMock, MagicMock

import openpyxl
from data.constants import ALL_GROUPS, ID_TO_NAME, get_kurs
from data.edupage_catalog import COURSE_ID_RANGES, catalog_groups, catalog_subjects, load_catalog
from keyboards.inline.jadval import get_kurs_keyboard
from utils.db_api.db import merge_edupage_groups
from utils.db_api.models import EdupageGroup, SubjectKafedra
from utils.excel_manager import (
    generate_kafedra_template, parse_kafedra_map_excel,
    generate_edupage_groups_template, parse_edupage_groups_excel,
)
from utils.scraper import _resolve_kafedra, scrape_one_group_with_driver, scrape_timetable


class CachedDriver:
    def quit(self):
        pass

    def get(self, url):
        self.current_url = url
        if url == 'about:blank':
            self.page_source = '<html></html>'
            return
        gid = url.split('class=*')[1]
        self.page_source = Path(f'data/edupage_94_pages/{gid}.html').read_text(encoding='utf-8')

    def find_elements(self, *args):
        return [] if self.current_url == 'about:blank' else [True]


class EdupageUpdateTests(unittest.TestCase):
    def test_combined_sheet_keeps_manual_columns_blank(self):
        from datetime import datetime
        from utils.sheets import upload_to_sheets
        client, service = MagicMock(), MagicMock()
        sheet = client.open_by_key.return_value.worksheet.return_value
        sheet.row_count = 3
        sheet._properties = {'sheetId': 1}
        lessons = [{'Filtr uchun': f'{k}-kurs', 'Guruh': f'Group{k}', 'Juft-lik': '1'} for k in range(1, 5)]
        with patch('utils.sheets._get_gspread_client', return_value=(client, service)):
            url, _ = upload_to_sheets(lessons, datetime(2026, 9, 7), 'new-sheet-id')
        client.open_by_key.assert_called_once_with('new-sheet-id')
        self.assertIn('new-sheet-id', url)
        self.assertEqual(service.spreadsheets().batchUpdate.call_args.kwargs['spreadsheetId'], 'new-sheet-id')
        rows = sheet.update.call_args.kwargs['values']
        self.assertEqual(rows[0][12], '2026-2027')
        self.assertEqual(rows[1][5], 'Kafedrasi')
        self.assertEqual(rows[3][0:2], ['1-kurs. I', '1-kurs, JAMI'])
        self.assertEqual([rows[2][i] for i in (8, 9, 10, 11, 14)], ['', '', '', '', ''])
        self.assertEqual([rows[3][i] for i in (8, 9, 10, 11, 14)], ['', '', '', '', ''])
        self.assertEqual(rows[-1][0:2], ['4-kurs. I', '4-kurs, JAMI'])
        self.assertEqual([rows[-1][i] for i in (8, 9, 10, 11, 14)], ['', '', '', '', ''])
        sheet.resize.assert_called_once_with(rows=13)

    def test_all_courses_reads_actual_pages_and_skips_empty_groups(self):
        from datetime import datetime
        with patch('utils.scraper.get_driver', return_value=CachedDriver()):
            lessons = scrape_timetable(datetime(2026, 9, 7), kurs='all')
        self.assertEqual({x['Filtr uchun'] for x in lessons}, {'1-kurs', '2-kurs', '3-kurs', '4-kurs'})
        self.assertTrue(all(x['Fan nomi'] for x in lessons))
        self.assertNotIn('PTI-76/23i', {x['Guruh'] for x in lessons})
        self.assertEqual([x['t/r'] for x in lessons], list(range(1, len(lessons) + 1)))
        callbacks = [b.callback_data for row in get_kurs_keyboard('2026-09-07').inline_keyboard for b in row]
        self.assertIn('kurs_all_2026-09-07', callbacks)

    def test_catalog_complete_and_courses_exact(self):
        self.assertEqual(load_catalog()['failures'], [])
        expected = {str(gid) for a, b in COURSE_ID_RANGES.values() for gid in range(a, b + 1)}
        self.assertEqual(set(ID_TO_NAME), expected)
        self.assertEqual(len(ALL_GROUPS), 132)
        for group in catalog_groups():
            self.assertEqual(get_kurs(group['id']), group['kurs'])
            self.assertIn('num=94&class=*' + group['id'], group['url'])
        for gid in ('253', '298', '304', '335', '340', '366', '371', '405', 'old'):
            self.assertEqual(get_kurs(gid), '')

    def test_current_catalog_not_hidden_by_old_database_rows(self):
        old = EdupageGroup(edupage_id='164', group_name='Old', kurs='1-kurs', num_param='90')
        override = EdupageGroup(edupage_id='254', group_name='Custom', kurs='1-kurs', num_param='94')
        merged = {g.edupage_id: g for g in merge_edupage_groups([old, override])}
        self.assertEqual(len(merged), 132)
        self.assertEqual(merged['254'].group_name, 'Custom')
        self.assertNotIn('164', merged)

    def test_every_course_uses_current_fallback_links(self):
        from datetime import datetime
        for course, (first, last) in COURSE_ID_RANGES.items():
            with patch('utils.scraper.get_driver') as driver, patch('utils.scraper.scrape_one_group_with_driver', return_value=[]) as scrape:
                scrape_timetable(datetime(2026, 9, 7), kurs=course)
                groups = [call.args[1] for call in scrape.call_args_list]
                self.assertEqual([g['id'] for g in groups], [str(g) for g in range(first, last + 1)])
                self.assertTrue(all('num=94&' in g['url'] and g['kurs'] == f'{course}-kurs' for g in groups))
                driver.return_value.quit.assert_called_once()
        buttons = [b.callback_data for row in get_kurs_keyboard('2026-09-07').inline_keyboard for b in row]
        for course in COURSE_ID_RANGES:
            self.assertIn(f'kurs_{course}_2026-09-07', buttons)

    def test_actual_pages_and_laboratory_normalization(self):
        catalog = {g['id']: g for g in catalog_groups()}
        for gid in ('254', '305', '341', '375'):
            lessons = scrape_one_group_with_driver(CachedDriver(), catalog[gid], 'Dushanba')
            self.assertTrue(lessons, gid)
            self.assertTrue(all(x['Guruh'] == catalog[gid]['name'] for x in lessons))
            self.assertTrue(all(x['Filtr uchun'] == catalog[gid]['kurs'] for x in lessons))
            self.assertTrue(all(x['Fan nomi'] in catalog_subjects() for x in lessons))
        lessons = scrape_one_group_with_driver(CachedDriver(), catalog['305'], 'Seshanba')
        self.assertTrue(any(x['Fan nomi'] == 'Ekonometrikaga kirish' and x["Mashg'ulot turi"] == 'Laboratoriya' for x in lessons))

    def test_template_complete_blank_and_roundtrip(self):
        data = generate_kafedra_template()
        wb = openpyxl.load_workbook(io.BytesIO(data))
        ws = wb.active
        subjects = [r[0] for r in ws.iter_rows(min_row=2, values_only=True)]
        self.assertEqual(subjects, catalog_subjects())
        self.assertEqual(len(subjects), 23)
        self.assertEqual(parse_kafedra_map_excel(data), [])
        ws['B2'] = 'Ijtimoiy fanlar'
        buf = io.BytesIO()
        wb.save(buf)
        self.assertEqual(parse_kafedra_map_excel(buf.getvalue()), [(subjects[0], 'Ijtimoiy fanlar')])

    def test_template_retains_existing_assignments(self):
        existing = [SubjectKafedra(subject_name='Statistika', kafedra_name='Statistika kafedrasi'),
                    SubjectKafedra(subject_name='Qo‘shimcha fan', kafedra_name='Ijtimoiy fanlar')]
        data = generate_kafedra_template(existing)
        self.assertEqual(set(parse_kafedra_map_excel(data)), {(s.subject_name, s.kafedra_name) for s in existing})

    def test_numbered_blank_rows_and_department_containing_fan(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(['№', 'Fan nomi', 'Tegishli Kafedra'])
        ws.append([1, 'Falsafa', None])
        ws.append([2, 'Mustaqillik darsi', 'Ijtimoiy fanlar kafedrasi'])
        buf = io.BytesIO()
        wb.save(buf)
        self.assertEqual(parse_kafedra_map_excel(buf.getvalue()), [('Mustaqillik darsi', 'Ijtimoiy fanlar kafedrasi')])

    def test_edupage_template_roundtrip(self):
        rows = parse_edupage_groups_excel(generate_edupage_groups_template())
        self.assertEqual(len(rows), 132)
        self.assertEqual({r[0] for r in rows}, set(ID_TO_NAME))
        self.assertTrue(all(r[3] == '94' for r in rows))

    def test_department_exact_match_precedes_prefix(self):
        self.assertEqual(_resolve_kafedra('Mikroiqtisodiyot 1', {'Mikroiqtisodiyot': 'General', 'Mikroiqtisodiyot 1': 'Specific'}), 'Specific')


class AutomaticUploadTests(unittest.IsolatedAsyncioTestCase):
    async def test_auto_upload_uses_all_courses_once(self):
        import main
        from datetime import datetime
        groups = merge_edupage_groups([])
        lessons = [{'Filtr uchun': f'{k}-kurs'} for k in range(1, 5)]
        with patch.object(main, 'datetime') as clock, \
             patch.object(main, 'get_kafedra_map_dict', new=AsyncMock(return_value={})), \
             patch.object(main, 'get_edupage_id_to_name_dict', new=AsyncMock(return_value={})), \
             patch.object(main, 'get_all_edupage_groups', new=AsyncMock(return_value=groups)), \
             patch.object(main, 'get_admin_ids', new=AsyncMock(return_value=[])), \
             patch.object(main, 'get_spreadsheet_id', new=AsyncMock(return_value='new-sheet-id')), \
             patch.object(main, 'daily_sheet_exists', return_value=False), \
             patch.object(main, 'scrape_timetable', return_value=lessons) as scrape, \
             patch.object(main, 'upload_to_sheets', return_value=('url', 'sheet')) as upload, \
             patch.object(main, 'ADMINS', []):
            clock.now.return_value = datetime(2026, 9, 7)
            await main.auto_upload_today()
            self.assertEqual(scrape.call_args.args[2], 'all')
            self.assertEqual(len(scrape.call_args.args[4]), 132)
            upload.assert_called_once_with(lessons, datetime(2026, 9, 7), 'new-sheet-id')

    async def test_auto_upload_preserves_existing_daily_sheet(self):
        import main
        from datetime import datetime
        with patch.object(main, 'datetime') as clock, \
             patch.object(main, 'get_admin_ids', new=AsyncMock(return_value=[])), \
             patch.object(main, 'get_spreadsheet_id', new=AsyncMock(return_value='new-sheet-id')), \
             patch.object(main, 'daily_sheet_exists', return_value=True) as exists, \
             patch.object(main, 'scrape_timetable') as scrape, \
             patch.object(main, 'upload_to_sheets') as upload, \
             patch.object(main, 'ADMINS', []):
            clock.now.return_value = datetime(2026, 9, 14)
            await main.auto_upload_today()
        exists.assert_called_once_with(datetime(2026, 9, 14), 'new-sheet-id')
        scrape.assert_not_called()
        upload.assert_not_called()


if __name__ == '__main__':
    unittest.main()
