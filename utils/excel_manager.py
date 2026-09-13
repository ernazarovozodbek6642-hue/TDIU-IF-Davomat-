"""
excel_manager.py — Excel import, export, shablonlar va hisobotlar
"""
import io
from datetime import date
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from utils.db_api.models import Attendance, AttendanceSession, Student, Tutor, SubjectKafedra, EdupageGroup
from data.constants import shorten_name
from data.edupage_catalog import TIMETABLE_NUM, catalog_groups, catalog_subjects


def _get_styles():
    thin = Side(style="thin")
    return {
        "header_fill": PatternFill("solid", fgColor="1F4E79"),
        "group_fill": PatternFill("solid", fgColor="2E75B6"),
        "para_fill": PatternFill("solid", fgColor="BDD7EE"),
        "absent_fill": PatternFill("solid", fgColor="FFE0E0"),
        "present_fill": PatternFill("solid", fgColor="E2EFDA"),
        "white_font": Font(name="Times New Roman", bold=True, color="FFFFFF", size=10),
        "dark_bold": Font(name="Times New Roman", bold=True, size=10),
        "normal_font": Font(name="Times New Roman", size=10),
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin)
    }


def generate_attendance_history_excel(
    att_date: date,
    paras: list[int],
    records: list[Attendance],
    all_present_sessions: set[tuple[str, int]],
    group_student_counts: dict[str, int]
) -> bytes:
    styles = _get_styles()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = att_date.strftime('%d.%m.%Y')

    def cell(row, col, value, font=None, fill=None, align=None):
        c = ws.cell(row=row, column=col, value=value)
        if font:  c.font = font
        if fill:  c.fill = fill
        if align: c.alignment = align
        c.border = styles["border"]
        return c

    # Sarlavha
    ws.merge_cells("A1:F1")
    para_str = ", ".join(f"{p}-para" for p in sorted(paras))
    c = ws.cell(row=1, column=1, value=f"Yoqlama tarixi — {att_date.strftime('%d.%m.%Y')} | {para_str}")
    c.font = styles["white_font"]
    c.fill = styles["header_fill"]
    c.alignment = styles["center"]
    c.border = styles["border"]

    headers = ["Guruh", "Para", "Talabalar soni", "Kelgan", "Kelmagan", "Kelmagan talabalar"]
    for ci, h in enumerate(headers, 1):
        cell(2, ci, h, font=styles["dark_bold"], fill=styles["para_fill"], align=styles["center"])

    by_gp = defaultdict(list)
    for r in records:
        by_gp[(r.group_name, r.para)].append(r.student_name)
    for gp in all_present_sessions:
        if gp not in by_gp:
            by_gp[gp] = []

    row_i = 3
    for (group, para) in sorted(by_gp.keys()):
        absent_list = by_gp[(group, para)]
        total = group_student_counts.get(group, 0)
        is_all_present = len(absent_list) == 0
        kelgan = total if is_all_present else max(0, total - len(absent_list))
        short_names = "✅ Hammasi kelgan" if is_all_present else "; ".join(shorten_name(n) for n in absent_list)
        row_fill = styles["present_fill"] if is_all_present else styles["group_fill"]

        cell(row_i, 1, group, font=styles["dark_bold"], fill=row_fill, align=styles["center"])
        cell(row_i, 2, para, font=styles["normal_font"], align=styles["center"])
        cell(row_i, 3, total, font=styles["normal_font"], align=styles["center"])
        cell(row_i, 4, kelgan, font=styles["normal_font"], align=styles["center"])
        cell(row_i, 5, 0 if is_all_present else len(absent_list), font=styles["normal_font"], align=styles["center"])
        cell(row_i, 6, short_names, font=styles["normal_font"], align=styles["center"])

        if not is_all_present and absent_list:
            ws.cell(row=row_i, column=5).fill = styles["absent_fill"]
        row_i += 1

    widths = [(1, 15), (2, 8), (3, 14), (4, 10), (5, 10), (6, 60)]
    for col, width in widths:
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 30

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ── Talabalar Export / Import ──────────────────────────────

def export_students_excel(students: list[Student]) -> bytes:
    styles = _get_styles()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Talabalar"

    headers = ["№", "Guruh", "Talaba F.I.Sh."]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = styles["white_font"]
        c.fill = styles["header_fill"]
        c.alignment = styles["center"]
        c.border = styles["border"]

    for idx, stu in enumerate(students, 1):
        r = idx + 1
        ws.cell(row=r, column=1, value=idx).alignment = styles["center"]
        ws.cell(row=r, column=2, value=stu.group_name).alignment = styles["center"]
        ws.cell(row=r, column=3, value=stu.full_name)
        for c in range(1, 4):
            ws.cell(row=r, column=c).border = styles["border"]
            ws.cell(row=r, column=c).font = styles["normal_font"]

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 45

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def parse_students_excel(file_bytes: bytes) -> list[tuple[str, str]]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    students = []

    for row in ws.iter_rows(values_only=True):
        if not row or all(v is None for v in row):
            continue
        vals = [str(v).strip() for v in row if v is not None]
        if not vals:
            continue

        if any("guruh" in v.lower() for v in vals) or any("talaba" in v.lower() for v in vals):
            continue

        if len(vals) >= 3 and vals[0].isdigit():
            group = vals[1]
            name = vals[2]
        elif len(vals) >= 2:
            group = vals[0]
            name = vals[1]
        else:
            continue

        if group and name and len(name) > 3:
            students.append((group, name))

    return students


def parse_single_group_students_excel(file_bytes: bytes, target_group: str) -> list[str]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    names = []

    for row in ws.iter_rows(values_only=True):
        if not row or all(v is None for v in row):
            continue
        vals = [str(v).strip() for v in row if v is not None]
        if not vals:
            continue

        if any("guruh" in v.lower() for v in vals) or any("talaba" in v.lower() for v in vals) or any("f.i.sh" in v.lower() for v in vals):
            continue

        name = ""
        if len(vals) == 1:
            name = vals[0]
        elif len(vals) >= 2 and vals[0].isdigit():
            name = vals[1]
        elif len(vals) >= 2:
            name = vals[1]

        if name and len(name) > 3:
            names.append(name)

    return names


# ── Tyutorlar Export / Import ──────────────────────────────

def export_tutors_excel(tutors_with_groups: list[tuple[Tutor, list[str]]]) -> bytes:
    styles = _get_styles()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tyutorlar"

    headers = ["№", "Telegram ID", "Tyutor F.I.Sh.", "Biriktirilgan guruhlar"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = styles["white_font"]
        c.fill = styles["header_fill"]
        c.alignment = styles["center"]
        c.border = styles["border"]

    for idx, (tutor, groups) in enumerate(tutors_with_groups, 1):
        r = idx + 1
        ws.cell(row=r, column=1, value=idx).alignment = styles["center"]
        ws.cell(row=r, column=2, value=tutor.telegram_id).alignment = styles["center"]
        ws.cell(row=r, column=3, value=tutor.name)
        ws.cell(row=r, column=4, value=", ".join(groups))
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = styles["border"]
            ws.cell(row=r, column=c).font = styles["normal_font"]

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 35
    ws.column_dimensions["D"].width = 40

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def parse_tutors_excel(file_bytes: bytes) -> list[tuple[int, str, list[str]]]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    tutors = []

    for row in ws.iter_rows(values_only=True):
        if not row or all(v is None for v in row):
            continue
        vals = [str(v).strip() for v in row if v is not None]
        if not vals:
            continue

        if any("telegram" in v.lower() for v in vals) or any("tyutor" in v.lower() for v in vals):
            continue

        tg_id_str = ""
        name = ""
        groups_str = ""

        if len(vals) >= 4 and vals[0].isdigit() and vals[1].isdigit():
            tg_id_str = vals[1]
            name = vals[2]
            groups_str = vals[3]
        elif len(vals) >= 3 and vals[0].isdigit():
            tg_id_str = vals[0]
            name = vals[1]
            groups_str = vals[2]

        if tg_id_str.isdigit() and name:
            tg_id = int(tg_id_str)
            grps = [g.strip() for g in groups_str.replace(";", ",").split(",") if g.strip()]
            tutors.append((tg_id, name, grps))

    return tutors


# ── Edupage Guruhlar Xaritasi Export / Import ──────────────

def export_edupage_groups_excel(groups: list[EdupageGroup]) -> bytes:
    """Edupage Guruhlari (ID, Guruh, Kurs) ro'yxatini Excel ga eksport qilish"""
    styles = _get_styles()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Edupage_Guruhlar"

    headers = ["EduPage ID", "Guruh nomi", "Kursi", "Timetable Num"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = styles["white_font"]
        c.fill = styles["header_fill"]
        c.alignment = styles["center"]
        c.border = styles["border"]

    for idx, grp in enumerate(groups, 2):
        ws.cell(row=idx, column=1, value=grp.edupage_id).alignment = styles["center"]
        ws.cell(row=idx, column=2, value=grp.group_name).alignment = styles["center"]
        ws.cell(row=idx, column=3, value=grp.kurs).alignment = styles["center"]
        ws.cell(row=idx, column=4, value=grp.num_param or TIMETABLE_NUM).alignment = styles["center"]
        for c in range(1, 5):
            ws.cell(row=idx, column=c).border = styles["border"]
            ws.cell(row=idx, column=c).font = styles["normal_font"]

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def parse_edupage_groups_excel(file_bytes: bytes) -> list[tuple[str, str, str, str]]:
    """
    Excel dan Edupage Guruhlarini o'qish
    Ustunlar: [Edupage ID | Guruh nomi | Kursi] yoki [Edupage ID | Guruh nomi | Kursi | Timetable Num]
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    items = []

    for row in ws.iter_rows(values_only=True):
        if not row or all(v is None for v in row):
            continue
        vals = [str(v).strip() for v in row if v is not None]
        if not vals:
            continue

        if any("id" in v.lower() for v in vals) or any("guruh" in v.lower() for v in vals):
            continue

        edupage_id = ""
        group_name = ""
        kurs = "1-kurs"
        num_param = TIMETABLE_NUM

        if len(vals) >= 4:
            edupage_id = vals[0]
            group_name = vals[1]
            kurs = vals[2]
            num_param = vals[3]
        elif len(vals) == 3:
            edupage_id = vals[0]
            group_name = vals[1]
            kurs = vals[2]
            num_param = TIMETABLE_NUM
        elif len(vals) == 2:
            edupage_id = vals[0]
            group_name = vals[1]
            num_param = TIMETABLE_NUM

        if edupage_id and group_name:
            items.append((edupage_id, group_name, kurs, num_param))

    return items


def generate_edupage_groups_template() -> bytes:
    """EduPage guruhlari importi uchun namunaviy shablon"""
    styles = _get_styles()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Edupage_Guruhlar_Shablon"

    headers = ["EduPage ID", "Guruh nomi", "Kursi", "Timetable Num (Ixtiyoriy)"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = styles["white_font"]
        c.fill = styles["header_fill"]
        c.alignment = styles["center"]

    sample_data = [(g['id'], g['name'], g['kurs'], TIMETABLE_NUM) for g in catalog_groups()]
    for r_idx, (gid, gname, kurs, num_p) in enumerate(sample_data, 2):
        ws.cell(row=r_idx, column=1, value=gid).alignment = styles["center"]
        ws.cell(row=r_idx, column=2, value=gname).alignment = styles["center"]
        ws.cell(row=r_idx, column=3, value=kurs).alignment = styles["center"]
        ws.cell(row=r_idx, column=4, value=num_p).alignment = styles["center"]

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 25

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ── Fanlar va Kafedralar Export / Import ───────────────────

def export_kafedra_map_excel(subject_kafedras: list[SubjectKafedra]) -> bytes:
    styles = _get_styles()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kafedralar"

    headers = ["№", "Fan nomi", "Tegishli Kafedra"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = styles["white_font"]
        c.fill = styles["header_fill"]
        c.alignment = styles["center"]
        c.border = styles["border"]

    for idx, item in enumerate(subject_kafedras, 1):
        r = idx + 1
        ws.cell(row=r, column=1, value=idx).alignment = styles["center"]
        ws.cell(row=r, column=2, value=item.subject_name)
        ws.cell(row=r, column=3, value=item.kafedra_name)
        for c in range(1, 4):
            ws.cell(row=r, column=c).border = styles["border"]
            ws.cell(row=r, column=c).font = styles["normal_font"]

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 40

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def parse_kafedra_map_excel(file_bytes: bytes) -> list[tuple[str, str]]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb.active
    items = []

    for row in ws.iter_rows(values_only=True):
        if not row or all(v is None for v in row):
            continue
        # Keep column positions: a blank department must never shift into a name.
        vals = [str(v).strip() if v is not None else '' for v in row]
        if not vals:
            continue

        if any(v.casefold() in ('fan nomi', 'tegishli kafedra') for v in vals):
            continue

        subj = ""
        kaf = ""
        if vals[0].isdigit():
            if len(vals) < 3:
                continue
            subj, kaf = vals[1:3]
        elif len(vals) >= 2:
            subj = vals[0]
            kaf = vals[1]

        if subj and kaf:
            items.append((subj, kaf))

    return items


# ── Boshqa Shablonlar ─────────────────────────────────────

def generate_student_template() -> bytes:
    styles = _get_styles()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Talabalar_Shablon"

    headers = ["Guruh", "Talaba F.I.Sh."]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = styles["white_font"]
        c.fill = styles["header_fill"]
        c.alignment = styles["center"]

    sample_data = [
        ("I-50/24", "ALIEV VALI AZIZOVICH"),
        ("I-50/24", "KARIMOV JASUR BOTIROVICH"),
        ("I-51/24", "SABIROVA MALIKA RUSTAM QIZI"),
    ]
    for r_idx, (g, n) in enumerate(sample_data, 2):
        ws.cell(row=r_idx, column=1, value=g).alignment = styles["center"]
        ws.cell(row=r_idx, column=2, value=n)

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 40

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def generate_single_group_template(group_name: str = "I-50/24") -> bytes:
    styles = _get_styles()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{group_name}_Talabalari"

    headers = ["№", "Talaba F.I.Sh."]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = styles["white_font"]
        c.fill = styles["header_fill"]
        c.alignment = styles["center"]

    sample_data = [
        "ALIEV VALI AZIZOVICH",
        "KARIMOV JASUR BOTIROVICH",
        "SABIROVA MALIKA RUSTAM QIZI",
    ]
    for r_idx, name in enumerate(sample_data, 2):
        ws.cell(row=r_idx, column=1, value=r_idx - 1).alignment = styles["center"]
        ws.cell(row=r_idx, column=2, value=name)

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 40

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def generate_tutor_template() -> bytes:
    styles = _get_styles()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tyutorlar_Shablon"

    headers = ["Telegram ID", "Tyutor F.I.Sh.", "Biriktirilgan guruhlar (vergul bilan)"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = styles["white_font"]
        c.fill = styles["header_fill"]
        c.alignment = styles["center"]

    sample_data = [
        ("5589013665", "Rahmonov Sadullo Shuxrat o'g'li", "I-50/24, I-51/24"),
        ("1234567890", "Safarova Sevara Shokir qizi", "I-52/24, I-53/24"),
    ]
    for r_idx, (tid, name, grps) in enumerate(sample_data, 2):
        ws.cell(row=r_idx, column=1, value=tid).alignment = styles["center"]
        ws.cell(row=r_idx, column=2, value=name)
        ws.cell(row=r_idx, column=3, value=grps)

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 45

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def generate_kafedra_template(existing: list[SubjectKafedra] | None = None) -> bytes:
    styles = _get_styles()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kafedralar_Shablon"

    headers = ["Fan nomi", "Tegishli Kafedra"]
    for ci, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.font = styles["white_font"]
        c.fill = styles["header_fill"]
        c.alignment = styles["center"]

    assignments = {item.subject_name: item.kafedra_name for item in (existing or [])}
    subjects = sorted(set(catalog_subjects()) | set(assignments), key=str.casefold)
    for r_idx, subj in enumerate(subjects, 2):
        ws.cell(row=r_idx, column=1, value=subj)
        ws.cell(row=r_idx, column=2, value=assignments.get(subj, ''))
        for col in (1, 2):
            ws.cell(row=r_idx, column=col).alignment = Alignment(vertical='center', wrap_text=True)
        ws.cell(row=r_idx, column=2).fill = PatternFill('solid', fgColor='FFF2CC')
        ws.row_dimensions[r_idx].height = 36

    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 45
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:B{len(subjects) + 1}'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
